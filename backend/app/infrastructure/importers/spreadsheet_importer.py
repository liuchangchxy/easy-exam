# Portions of the column mapping and spreadsheet parsing algorithms below
# are adapted from Exameow (frontend/src/utils/importParser.ts at commit 70e0d70),
# licensed under the Apache License, Version 2.0.
# Copyright (c) Exameow contributors.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
#
# A complete copy of the Apache License, Version 2.0 is provided in
# licenses/APACHE-2.0.txt and backend/licenses/APACHE-2.0.txt.

import csv
import io
import json
import re
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook


def _normalize(s: str) -> str:
    return re.sub(r"\s+", "", str(s or "").strip().lower())


def normalize_difficulty(value: Any) -> int:
    """Normalize textual or numeric difficulty into EasyExam 1..5 scale (default 3)."""
    norm = _normalize(str(value or ""))
    if norm in {"简单", "easy", "1"}:
        return 1
    if norm in {"中等", "适中", "medium", "2", "3"}:
        return 3
    if norm in {"困难", "hard", "4", "5"}:
        return 5
    try:
        val = int(value)
        return min(max(val, 1), 5)
    except (ValueError, TypeError):
        return 3


def detect_column_type(val: str) -> Optional[str]:
    v = _normalize(val)
    if "单选" in v or ("single" in v and "multi" not in v):
        return "SINGLE"
    if "多选" in v or "multi" in v or "多项" in v:
        return "MULTI"
    if "判断" in v or "true" in v or "false" in v or "对错" in v or "是非" in v:
        return "JUDGE"
    if "填空" in v or "fill" in v or "blank" in v:
        return "FILL"
    if "简答" in v or "问答" in v or "short" in v or "essay" in v or "主观" in v:
        return "ESSAY"
    return None


def detect_type_from_qa(stem: str, answer: str, has_options: bool) -> str:
    ans_clean = str(answer or "").strip()
    if has_options:
        if len(ans_clean) <= 4 and any(ch in ans_clean for ch in [",", ";", "、", "，", "；"]):
            return "MULTI"
        if len(ans_clean) > 1 and re.match(r"^[A-Ha-h]+$", ans_clean):
            return "MULTI"
        return "SINGLE"
    judge_keys = {"对", "错", "√", "×", "正确", "错误", "true", "false", "是", "否"}
    if ans_clean.lower().rstrip(".。") in judge_keys:
        return "JUDGE"
    if "____" in stem or "填空" in stem:
        return "FILL"
    return "SINGLE" if has_options else "FILL"


OPTION_PREFIX_SPLIT = re.compile(r"\s*[A-Ha-h][.、．:：)）]\s*")
OPTION_PREFIX_STRIP = re.compile(r"^\s*[A-Ha-h][.、．:：)）]\s*")


def strip_option_prefix(s: str) -> str:
    return re.sub(r"[;；|、]\s*$", "", OPTION_PREFIX_STRIP.sub("", s)).strip()


def split_with_delimiter(cell: str, delimiter: str) -> List[str]:
    parts = cell.split(delimiter)
    return [strip_option_prefix(p) for p in parts if strip_option_prefix(p)]


def split_options_cell(cell: str, delimiter: Optional[str] = None) -> List[str]:
    text = str(cell or "").strip()
    if not text:
        return []
    if delimiter == "prefix":
        parts = [p.strip() for p in OPTION_PREFIX_SPLIT.split(text) if p.strip()]
        return parts if parts else [text]
    if delimiter:
        sep = "\n" if delimiter == "\\n" else delimiter
        return split_with_delimiter(text, sep)

    if "\n" in text:
        res = split_with_delimiter(text, "\n")
        if len(res) >= 2:
            return res
    for d in ["；", ";", "|", "、"]:
        if d in text:
            res = split_with_delimiter(text, d)
            if len(res) >= 2:
                return res
    if OPTION_PREFIX_STRIP.search(text):
        parts = [p.strip() for p in OPTION_PREFIX_SPLIT.split(text) if p.strip()]
        if len(parts) >= 2:
            return parts
    return [text]


def detect_options_delimiter(cells: List[str]) -> str:
    for cell in cells:
        text = str(cell or "").strip()
        if not text:
            continue
        if "\n" in text and len(split_with_delimiter(text, "\n")) >= 2:
            return "\\n"
        for d in ["；", ";", "|", "、"]:
            if d in text and len(split_with_delimiter(text, d)) >= 2:
                return d
        if OPTION_PREFIX_STRIP.search(text):
            parts = [p.strip() for p in OPTION_PREFIX_SPLIT.split(text) if p.strip()]
            if len(parts) >= 2:
                return "prefix"
    return ""


def build_column_map(headers: List[str]) -> Dict[str, Any]:
    mapping: Dict[str, Any] = {
        "stem": None,
        "type": None,
        "options": [],
        "combined_options": None,
        "options_delimiter": "",
        "answer": None,
        "explanation": None,
        "difficulty": None,
        "tags": None,
    }
    letters = "ABCDEFGH"
    slots: List[Optional[int]] = [None] * 8
    next_unnamed_slot = 0

    for i, h in enumerate(headers):
        if not h:
            continue
        n = _normalize(h)
        if mapping["stem"] is None and (
            "题干" in n or "题目" in n or n == "题" or "stem" in n or "question" in n or n == "title" or "内容" in n or n == "q"
        ):
            mapping["stem"] = i
            continue
        if mapping["type"] is None and (
            "题型" in n or "类型" in n or "type" in n or n == "qt" or "种类" in n
        ):
            mapping["type"] = i
            continue
        if mapping["answer"] is None and (
            "答案" in n or "正确" in n or "answer" in n or n == "ans" or "标准" in n or "key" in n
        ):
            mapping["answer"] = i
            continue
        if mapping["explanation"] is None and (
            "解析" in n or "分析" in n or "analysis" in n or "explanation" in n or "详解" in n or "解释" in n
        ):
            mapping["explanation"] = i
            continue
        if mapping["tags"] is None and (
            "章节" in n or "chapter" in n or "unit" in n or "模块" in n or
            "学科" in n or "科目" in n or "课程" in n or "subject" in n or
            "标签" in n or "考点" in n or "分类" in n
        ):
            mapping["tags"] = i
            continue
        if mapping["difficulty"] is None and (
            "难度" in n or "難度" in n or "難易度" in n or "难易度" in n or "difficulty" in n
        ):
            mapping["difficulty"] = i
            continue
        if n in {"选项", "options", "option", "所有选项", "全部选项", "选项内容", "choices", "choice"}:
            if mapping["combined_options"] is None:
                mapping["combined_options"] = i
            continue

        slot = None
        if len(n) == 1 and n.upper() in letters:
            slot = letters.index(n.upper())
        else:
            m = re.search(r"(?:选项|option)\s*([a-hA-H])(?:\b|$|\s|列)", h)
            if m:
                slot = letters.index(m.group(1).upper())
        if slot is not None:
            slots[slot] = i
            continue
        if "选项" in n or "option" in n:
            while next_unnamed_slot < 8 and slots[next_unnamed_slot] is not None:
                next_unnamed_slot += 1
            if next_unnamed_slot < 8:
                slots[next_unnamed_slot] = i
                next_unnamed_slot += 1
            continue

    while slots and slots[-1] is None:
        slots.pop()
    mapping["options"] = slots
    return mapping


def analyze_spreadsheet_structure(headers: List[str], sample_rows: List[List[Any]]) -> Dict[str, Any]:
    mapping = build_column_map(headers)
    if mapping["combined_options"] is not None:
        idx = mapping["combined_options"]
        samples = [str(r[idx]) for r in sample_rows[:20] if idx < len(r) and r[idx] is not None]
        mapping["options_delimiter"] = detect_options_delimiter(samples)
    missing = []
    if mapping["stem"] is None:
        missing.append("stem")
    if mapping["answer"] is None:
        missing.append("answer")
    return {
        "headers": headers,
        "mapping": mapping,
        "missing": missing,
    }


def parse_rows_with_mapping(rows: List[List[Any]], mapping: Dict[str, Any]) -> List[Dict[str, Any]]:
    stem_idx = mapping.get("stem")
    ans_idx = mapping.get("answer")
    type_idx = mapping.get("type")
    exp_idx = mapping.get("explanation")
    diff_idx = mapping.get("difficulty")
    tags_idx = mapping.get("tags")
    opt_idxs = mapping.get("options") or []
    comb_idx = mapping.get("combined_options")
    delimiter = mapping.get("options_delimiter") or None

    items = []
    for row in rows:
        if not row or not any(str(c or "").strip() for c in row):
            continue
        stem = str(row[stem_idx]).strip() if stem_idx is not None and stem_idx < len(row) and row[stem_idx] is not None else ""
        if not stem:
            continue
        answer = str(row[ans_idx]).strip() if ans_idx is not None and ans_idx < len(row) and row[ans_idx] is not None else ""
        qtype_raw = str(row[type_idx]).strip() if type_idx is not None and type_idx < len(row) and row[type_idx] is not None else ""
        qtype = detect_column_type(qtype_raw) if qtype_raw else None

        options = []
        if comb_idx is not None and comb_idx < len(row) and row[comb_idx] is not None:
            cell_val = row[comb_idx]
            if isinstance(cell_val, list):
                options = cell_val
            elif isinstance(cell_val, str) and cell_val.strip().startswith("["):
                try:
                    options = json.loads(cell_val)
                except Exception:
                    raw_opts = split_options_cell(cell_val, delimiter)
                    letters = "ABCDEFGH"
                    options = [{"key": letters[idx] if idx < len(letters) else str(idx + 1), "content": opt} for idx, opt in enumerate(raw_opts)]
            else:
                raw_opts = split_options_cell(str(cell_val), delimiter)
                letters = "ABCDEFGH"
                options = [{"key": letters[idx] if idx < len(letters) else str(idx + 1), "content": opt} for idx, opt in enumerate(raw_opts)]
        elif opt_idxs:
            letters = "ABCDEFGH"
            if isinstance(opt_idxs, dict):
                def _sort_key(item):
                    k = str(item[0]).strip().upper()
                    return letters.index(k) if k in letters else (int(k) if k.isdigit() else 99)
                for k, col in sorted(opt_idxs.items(), key=_sort_key):
                    if col is not None and str(col).strip().isdigit():
                        c_idx = int(col)
                        if c_idx < len(row) and row[c_idx] is not None and str(row[c_idx]).strip():
                            key = k.upper() if k.upper() in letters else (letters[int(k)] if k.isdigit() and int(k) < len(letters) else str(k))
                            options.append({"key": key, "content": strip_option_prefix(str(row[c_idx]).strip())})
            elif isinstance(opt_idxs, list):
                for oi_idx, oi in enumerate(opt_idxs):
                    if oi is None:
                        continue
                    if isinstance(oi, dict):
                        key = oi.get("key", letters[oi_idx] if oi_idx < len(letters) else str(oi_idx + 1))
                        col = oi.get("column", oi.get("col"))
                        if col is not None and str(col).strip().isdigit():
                            c_idx = int(col)
                            if c_idx < len(row) and row[c_idx] is not None and str(row[c_idx]).strip():
                                options.append({"key": str(key).upper(), "content": strip_option_prefix(str(row[c_idx]).strip())})
                    elif isinstance(oi, (int, float, str)) and str(oi).strip().isdigit():
                        col = int(oi)
                        if col < len(row) and row[col] is not None and str(row[col]).strip():
                            key = letters[oi_idx] if oi_idx < len(letters) else str(oi_idx + 1)
                            options.append({"key": key, "content": strip_option_prefix(str(row[col]).strip())})

        if not qtype:
            qtype = detect_type_from_qa(stem, answer, bool(options))

        explanation = str(row[exp_idx]).strip() if exp_idx is not None and exp_idx < len(row) and row[exp_idx] is not None else ""
        difficulty = normalize_difficulty(row[diff_idx]) if diff_idx is not None and diff_idx < len(row) and row[diff_idx] is not None else 3

        tags = []
        if tags_idx is not None and tags_idx < len(row) and row[tags_idx] is not None:
            tag_val = str(row[tags_idx]).replace("，", ",").strip()
            tags = [p.strip() for p in tag_val.split(",") if p.strip()]

        items.append({
            "stem": stem,
            "answer": answer,
            "type": qtype or "SINGLE",
            "explanation": explanation,
            "difficulty": difficulty,
            "tags": tags,
            "options": options,
        })
    return items


def read_xlsx_rows(content: bytes) -> Tuple[List[str], List[List[Any]]]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        raw_rows = list(sheet.iter_rows(values_only=True))
    except Exception as exc:
        raise ValueError("Excel 文件无法读取或已损坏") from exc
    if not raw_rows:
        raise ValueError("Excel 文件为空")
    headers = [str(c or "").strip() for c in raw_rows[0]]
    rows = [list(r) for r in raw_rows[1:]]
    return headers, rows


def read_csv_rows(content: str) -> Tuple[List[str], List[List[Any]]]:
    try:
        reader = csv.reader(io.StringIO(content.lstrip("\ufeff")))
        all_rows = list(reader)
    except Exception as exc:
        raise ValueError(f"CSV 解析失败：{exc}") from exc
    if not all_rows:
        raise ValueError("CSV 文件为空")
    headers = [str(c or "").strip() for c in all_rows[0]]
    rows = [[str(c or "").strip() for c in r] for r in all_rows[1:]]
    return headers, rows


def parse_xlsx_questions(content: bytes, custom_mapping: Optional[Dict[str, Any]] = None) -> List[Dict[str, Any]]:
    headers, rows = read_xlsx_rows(content)
    analysis = analyze_spreadsheet_structure(headers, rows)
    mapping = custom_mapping or analysis["mapping"]
    if mapping.get("stem") is None or mapping.get("answer") is None:
        raise ValueError("Excel 至少需要题干和答案列")
    items = parse_rows_with_mapping(rows, mapping)
    if not items:
        raise ValueError("Excel 中未识别到有效题目")
    return items
